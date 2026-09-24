"""The config-driven adapter, for sources that are "just tables".

A spreadsheet export, a CSV drop, a read-only view on somebody's database —
anything with rows and columns. Instead of code per customer, there is a YAML
mapping per customer, and the code is written once.

This is the adapter that makes onboarding a messy export an afternoon rather
than a week, and `python -m app.onboard draft-mapping` writes the first draft of
that YAML for a person to review.

What it refuses to do is invent. A column that is not there produces a capability
that is false, which produces a tool the agent never sees, which produces "your
system doesn't record that" instead of a number.
"""

from __future__ import annotations

import logging
import re
from collections import Counter, defaultdict
from collections.abc import AsyncIterator
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.canonical.enums import Channel, OrderStatus
from app.canonical.models import (
    CanonicalCategory,
    CanonicalInventoryLevel,
    CanonicalLocation,
    CanonicalOrder,
    CanonicalOrderLine,
    CanonicalProduct,
    CanonicalVariant,
    Capabilities,
)
from app.config import REPO_ROOT
from app.connectors.adapters.mapping_spec import (
    EntitySpec,
    MappingError,
    MappingSpec,
    apply_transforms,
    normalise_header,
)
from app.connectors.adapters.mapping_spec import load as load_spec
from app.connectors.base import AdapterError, AdapterInfo, Fetched, HealthStatus, SourceAdapter
from app.connectors.registry import register

log = logging.getLogger(__name__)

ADAPTER_NAME = "mapping"


@dataclass
class Row:
    sheet: str
    number: int
    cells: dict[str, Any] = field(default_factory=dict)
    raw: dict[str, Any] = field(default_factory=dict)

    def get(self, header: str) -> Any:
        return self.cells.get(normalise_header(header))

    @property
    def is_blank(self) -> bool:
        return all(value in (None, "") for value in self.cells.values())


def jsonable(value: Any) -> Any:
    """`raw_records` holds jsonb, and a spreadsheet hands back datetimes."""
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


class Workbook:
    """The file, read once, as rows keyed by normalised header."""

    def __init__(self, path: Path, spec: MappingSpec) -> None:
        self.path = path
        self.spec = spec
        self._sheets: dict[str, list[Row]] = {}
        self._lookups: dict[tuple[str, str, str], dict[str, Any]] = {}

    def load(self) -> None:
        if self._sheets:
            return
        if not self.path.exists():
            raise AdapterError(f"the export is not where the mapping says it is: {self.path}")
        book = load_workbook(self.path, read_only=True, data_only=True)
        try:
            for key, sheet_spec in self.spec.sheets.items():
                pattern = re.compile(sheet_spec.match)
                matched = [name for name in book.sheetnames if pattern.match(name)]
                if not matched:
                    raise AdapterError(
                        f"no sheet matches {sheet_spec.match!r} for {key!r}; "
                        f"the file has {book.sheetnames}"
                    )
                rows: list[Row] = []
                for name in matched:
                    rows.extend(self._read(book[name], sheet_spec.header_row))
                self._sheets[key] = rows
                log.info("%s: %s rows from %s", key, len(rows), ", ".join(matched))
        finally:
            book.close()

    def _read(self, sheet: Any, header_row: int) -> list[Row]:
        rows: list[Row] = []
        headers: list[str] = []
        originals: list[str] = []
        skip = self.spec.skip_rows_matching

        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index < header_row:
                continue
            if index == header_row:
                originals = [
                    str(v) if v is not None else f"column_{i}" for i, v in enumerate(values)
                ]
                headers = [normalise_header(v) for v in values]
                continue

            row = Row(
                sheet=sheet.title,
                number=index,
                cells={h: v for h, v in zip(headers, values, strict=False) if h},
                raw={o: jsonable(v) for o, v in zip(originals, values, strict=False)},
            )
            if row.is_blank:
                # Blank separator rows are decoration, not data.
                continue
            texts = [str(v).strip().casefold() for v in row.cells.values() if v not in (None, "")]
            if skip and any(text in skip for text in texts):
                # The "TOTAL" line at the bottom of every export ever produced.
                # Counting it would double the shop's revenue.
                log.debug("skipping summary row %s:%s", row.sheet, row.number)
                continue
            rows.append(row)
        return rows

    def rows(self, key: str) -> list[Row]:
        self.load()
        if key not in self._sheets:
            raise MappingError(f"no sheet registered under {key!r}")
        return self._sheets[key]

    def lookup_table(
        self, sheet: str, key_column: str, value_column: str, key_transforms: list[Any]
    ) -> dict[str, Any]:
        cache_key = (sheet, normalise_header(key_column), normalise_header(value_column))
        if cache_key not in self._lookups:
            table: dict[str, Any] = {}
            for row in self.rows(sheet):
                key = apply_transforms(
                    row.get(key_column), key_transforms or ["fold_key"], f"lookup key {key_column}"
                )
                if key in (None, ""):
                    continue
                table.setdefault(str(key), row.get(value_column))
            self._lookups[cache_key] = table
        return self._lookups[cache_key]


class FieldResolver:
    def __init__(self, workbook: Workbook, spec: MappingSpec) -> None:
        self.workbook = workbook
        self.spec = spec

    def value(self, definition: Any, row: Row, where: str, group_key: str | None = None) -> Any:
        if not isinstance(definition, dict):
            # Bare string shorthand: the column name.
            definition = {"column": definition}

        transforms = definition.get("transforms") or []

        if "group_key" in definition:
            # The identity of the order this row belongs to. For a source with
            # no ids of its own this is the only value guaranteed unique, so
            # line ids have to be built on it rather than on a column that may
            # be blank.
            raw = group_key
        elif "constant" in definition:
            raw = definition["constant"]
        elif "column" in definition:
            raw = row.get(definition["column"])
        elif "row_number" in definition:
            raw = f"{row.sheet}:{row.number}"
        elif "concat" in definition:
            parts = [
                self.value(part, row, f"{where}.concat", group_key) for part in definition["concat"]
            ]
            separator = definition.get("separator", "")
            raw = separator.join("" if p is None else str(p) for p in parts)
        elif "lookup" in definition:
            options = definition["lookup"]
            table = self.workbook.lookup_table(
                options["sheet"],
                options["key"],
                options["value"],
                options.get("key_transforms") or ["fold_key"],
            )
            probe = apply_transforms(
                row.get(options["match_on"]),
                options.get("match_transforms") or ["fold_key"],
                f"{where}.lookup",
            )
            raw = table.get(str(probe)) if probe not in (None, "") else None
        else:
            raise MappingError(
                f"{where}: no source. Expected one of column, constant, concat, "
                f"lookup, row_number or group_key"
            )

        # `at_time` needs the tenant's timezone, which lives in settings rather
        # than being repeated on every date field.
        steps = [
            {"at_time": {"at": self.spec.assumed_time, "timezone": self.spec.timezone}}
            if step == "at_time"
            else step
            for step in transforms
        ]
        return apply_transforms(raw, steps, where)

    def record(
        self, fields: dict[str, Any], row: Row, where: str, group_key: str | None = None
    ) -> dict[str, Any]:
        return {
            name: self.value(definition, row, f"{where}.{name}", group_key)
            for name, definition in fields.items()
        }


def modal_row(rows: list[Row]) -> Row:
    """One row standing for a group, taking the most common value per column.

    A shop types the same item four ways across a year. The most common
    spelling is the one to show, and blanks never win — which also quietly
    fixes the 9% of rows where somebody left the category empty.
    """
    first = rows[0]
    merged = Row(sheet=first.sheet, number=first.number, raw=first.raw)
    columns = {header for row in rows for header in row.cells}
    for header in columns:
        counts = Counter(
            str(row.cells.get(header)) for row in rows if row.cells.get(header) not in (None, "")
        )
        if not counts:
            merged.cells[header] = None
            continue
        best = counts.most_common(1)[0][0]
        # Hand back the original typed value, not its string form.
        for row in rows:
            if str(row.cells.get(header)) == best:
                merged.cells[header] = row.cells.get(header)
                break
    return merged


class MappingAdapter(SourceAdapter):
    def __init__(self, config: dict[str, Any], secret: str | None = None) -> None:
        mapping_file = config.get("mapping_file")
        if not mapping_file:
            raise AdapterError(
                "the integration's config needs a `mapping_file` pointing at the YAML mapping"
            )
        self.spec = load_spec(mapping_file, REPO_ROOT)
        workbook_path = Path(config["workbook"]) if config.get("workbook") else self.spec.workbook
        if not workbook_path.is_absolute():
            workbook_path = REPO_ROOT / workbook_path
        self.workbook = Workbook(workbook_path, self.spec)
        self.resolver = FieldResolver(self.workbook, self.spec)

    # -- description --------------------------------------------------------

    def describe(self) -> AdapterInfo:
        return AdapterInfo(
            name=ADAPTER_NAME,
            display_name=self.spec.display_name,
            capabilities=Capabilities(**self.spec.capabilities),
            notes=tuple(self.spec.notes),
        )

    async def healthcheck(self) -> HealthStatus:
        try:
            self.workbook.load()
        except (AdapterError, MappingError) as exc:
            return HealthStatus(ok=False, detail=str(exc))
        counts = ", ".join(
            f"{key}: {len(self.workbook.rows(key))} rows" for key in self.spec.sheets
        )
        return HealthStatus(ok=True, detail=f"{self.workbook.path.name} — {counts}")

    # -- helpers ------------------------------------------------------------

    def _entity(self, name: str) -> EntitySpec | None:
        return self.spec.entities.get(name)

    def _usable_rows(self, entity: EntitySpec) -> list[Row]:
        rows = self.workbook.rows(entity.source_sheet) if entity.source_sheet else []
        if not entity.skip_when_blank:
            return rows
        return [
            row
            for row in rows
            if all(row.get(column) not in (None, "") for column in entity.skip_when_blank)
        ]

    def _grouped(self, entity: EntitySpec, definition: Any) -> dict[str, list[Row]]:
        groups: dict[str, list[Row]] = defaultdict(list)
        for row in self._usable_rows(entity):
            key = self.resolver.value(definition, row, f"{entity.name}.key")
            if key in (None, ""):
                # No receipt number: the row stands on its own rather than
                # being merged with every other unnumbered sale.
                key = f"{row.sheet}:{row.number}"
            groups[str(key)].append(row)
        return groups

    # -- entities -----------------------------------------------------------

    async def iter_locations(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalLocation]]:
        entity = self._entity("locations")
        if entity is None:
            return
        for constant in entity.constants:
            yield Fetched(
                record=CanonicalLocation(
                    external_id=str(constant["external_id"]),
                    name=constant["name"],
                    timezone=constant.get("timezone", self.spec.timezone),
                ),
                raw=dict(constant),
            )

    async def iter_categories(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalCategory]]:
        async for item in self._distinct("categories", CanonicalCategory):
            yield item

    async def iter_products(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalProduct]]:
        async for item in self._distinct("products", CanonicalProduct):
            yield item

    async def iter_variants(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalVariant]]:
        async for item in self._distinct("variants", CanonicalVariant):
            yield item

    async def _distinct(self, name: str, model: type) -> AsyncIterator[Fetched[Any]]:
        entity = self._entity(name)
        if entity is None or entity.distinct_by is None:
            return
        for key, rows in self._grouped(entity, entity.distinct_by).items():
            merged = modal_row(rows)
            values = self.resolver.record(entity.fields, merged, name)
            if not values.get("external_id"):
                continue
            yield Fetched(
                record=model(**{k: v for k, v in values.items() if v is not None}),
                raw={"_key": key, "_rows": len(rows), **merged.raw},
            )

    async def iter_orders(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalOrder]]:
        entity = self._entity("orders")
        if entity is None or entity.group_by is None:
            return
        line_fields = (entity.lines or {}).get("fields") or {}

        for key, rows in self._grouped(entity, entity.group_by).items():
            header = self.resolver.record(entity.fields, modal_row(rows), "orders", key)
            lines: list[CanonicalOrderLine] = []
            for index, row in enumerate(rows, start=1):
                values = self.resolver.record(line_fields, row, "orders.lines", key)
                values.setdefault("external_id", f"{key}:{index}")
                if not values.get("external_id"):
                    values["external_id"] = f"{key}:{index}"
                values["quantity"] = values.get("quantity") or Decimal("0")
                values["unit_price"] = values.get("unit_price") or Decimal("0")
                values["discount"] = values.get("discount") or Decimal("0")
                lines.append(
                    CanonicalOrderLine(**{k: v for k, v in values.items() if v is not None})
                )
            if not lines:
                continue

            # A table source has no order-level totals: they are the sum of the
            # rows, which is also what makes them reconcile by construction.
            subtotal = sum((line.quantity * line.unit_price for line in lines), Decimal("0"))
            discounts = sum((line.discount for line in lines), Decimal("0"))
            tax = header.get("tax_total") or Decimal("0")
            tip = header.get("tip_total") or Decimal("0")

            placed_at = header.get("placed_at")
            if placed_at is None:
                continue

            yield Fetched(
                record=CanonicalOrder(
                    external_id=str(header.get("external_id") or key),
                    location_external_id=header.get("location_external_id"),
                    customer_external_id=header.get("customer_external_id"),
                    status=OrderStatus(header.get("status") or OrderStatus.COMPLETED),
                    channel=Channel(header.get("channel") or Channel.IN_STORE),
                    subtotal=subtotal,
                    discount_total=discounts,
                    tax_total=tax,
                    tip_total=tip,
                    total=subtotal - discounts + tax + tip,
                    placed_at=placed_at,
                    source_updated_at=placed_at,
                    lines=lines,
                ),
                raw={"_key": key, "_rows": [row.raw for row in rows]},
            )

    async def iter_inventory_levels(
        self, since: datetime | None = None, resume_cursor: str | None = None
    ) -> AsyncIterator[Fetched[CanonicalInventoryLevel]]:
        entity = self._entity("inventory_levels")
        if entity is None:
            return
        as_of = datetime.now(tz=UTC)
        for row in self._usable_rows(entity):
            values = self.resolver.record(entity.fields, row, "inventory_levels")
            if not values.get("variant_external_id"):
                continue
            on_hand = values.get("on_hand")
            if on_hand is None:
                continue
            yield Fetched(
                record=CanonicalInventoryLevel(
                    external_id=(
                        f"{values['variant_external_id']}:{values.get('location_external_id')}"
                    ),
                    variant_external_id=str(values["variant_external_id"]),
                    location_external_id=str(values.get("location_external_id")),
                    on_hand=Decimal(on_hand),
                    # The file has no timestamp, so "as of" is when we read it.
                    as_of=values.get("as_of") or as_of,
                ),
                raw=row.raw,
            )


def _build(config: dict[str, Any], secret: str | None) -> SourceAdapter:
    return MappingAdapter(config, secret)


def _describe() -> AdapterInfo:
    # Without a mapping file there is nothing to promise: this adapter's
    # capabilities are whatever the customer's YAML declares.
    return AdapterInfo(
        name=ADAPTER_NAME,
        display_name="Config-driven mapping (spreadsheets, CSV, SQL views)",
        capabilities=Capabilities(),
        notes=("Capabilities come from the tenant's mapping file, not from this adapter.",),
    )


register(ADAPTER_NAME, _build, _describe)
