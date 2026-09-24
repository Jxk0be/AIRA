"""CSV and XLSX, for the person who is going to do sums with it.

A bookkeeper does not want a PDF of a table, they want the table. Both formats
are produced from the same `Sheet` so the two can never disagree about what a
column means, and numbers are written as numbers — a spreadsheet where every
total is text because it came through with a dollar sign is worse than no
spreadsheet.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

Cell = str | int | float | Decimal | date | datetime | None


@dataclass(slots=True)
class Sheet:
    """One table, named."""

    name: str
    columns: list[str]
    rows: list[list[Cell]] = field(default_factory=list)
    # Printed above the table in XLSX and as a leading comment line in CSV.
    note: str | None = None


def csv_bytes(sheet: Sheet) -> bytes:
    """One sheet as CSV, UTF-8 with a BOM.

    The BOM is there because the overwhelmingly likely next step is
    double-clicking the file on a Windows machine, and Excel without it turns
    every accented product name into mojibake.
    """
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    if sheet.note:
        writer.writerow([f"# {sheet.note}"])
    writer.writerow(sheet.columns)
    for row in sheet.rows:
        writer.writerow([_csv_value(value) for value in row])
    return buffer.getvalue().encode("utf-8-sig")


def _csv_value(value: Cell) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)


def workbook_bytes(sheets: list[Sheet], *, title: str | None = None) -> bytes:
    """Several sheets in one workbook, one tab each."""
    book = Workbook()
    book.remove(book.active)
    if title:
        book.properties.title = title

    for sheet in sheets:
        # Excel refuses these characters in a tab name and silently truncates
        # past 31, which turns two long tab names into one collision.
        tab = book.create_sheet(_safe_tab_name(sheet.name))
        row_index = 1
        if sheet.note:
            tab.cell(row=1, column=1, value=sheet.note).font = Font(italic=True, color="666666")
            row_index = 3

        for column_index, heading in enumerate(sheet.columns, start=1):
            cell = tab.cell(row=row_index, column=column_index, value=heading)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="left")

        for offset, row in enumerate(sheet.rows, start=row_index + 1):
            for column_index, value in enumerate(row, start=1):
                tab.cell(row=offset, column=column_index, value=_xlsx_value(value))

        tab.freeze_panes = tab.cell(row=row_index + 1, column=1)
        _autosize(tab, sheet)

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _xlsx_value(value: Cell) -> Any:
    """Numbers as numbers, dates as dates, so the sheet can be worked with."""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _safe_tab_name(name: str) -> str:
    cleaned = "".join("-" if char in "[]:*?/\\" else char for char in name)
    return cleaned[:31] or "Sheet"


def _autosize(tab: Any, sheet: Sheet) -> None:
    for index, heading in enumerate(sheet.columns, start=1):
        widest = len(str(heading))
        for row in sheet.rows[:200]:
            if index <= len(row):
                widest = max(widest, len(_csv_value(row[index - 1])))
        tab.column_dimensions[get_column_letter(index)].width = min(widest + 3, 48)
