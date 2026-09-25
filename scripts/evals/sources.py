"""Ground truth, taken from the customer's own system.

Every expected answer in the golden sets is computed here, from the source the
shop actually runs — RegisterOne's database for Animanga Knox, the raw spreadsheet
for Panel & Pawn. Never from our canonical copy.

That is the only way an eval can catch a mistake in the adapter. If the truth
came from our own tables, an adapter that dropped every con-booth sale would
agree with itself perfectly and the eval would be green.

The spreadsheet side re-implements the parsing — money with dollar signs, dates
as "3/7/25", names typed three different ways — rather than reusing the mapping
adapter's. Two implementations that agree are evidence; one implementation
checked against itself is not.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

# Rows the export ends each sheet with. They are not sales, and counting them
# roughly doubles the shop's revenue.
SUMMARY_ROWS = ("total", "subtotal", "grand total")

DATE_FORMATS = ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d")


class OracleError(RuntimeError):
    """The ground truth could not be computed. Never guessed at."""


def fold(name: Any) -> str:
    """The same name however it was typed.

    " Iron Gutter #3", "IRON GUTTER #3" and "Iron  Gutter #3" are one item, and
    the mapping adapter folds them the same way — independently, which is the
    point of writing it twice.
    """
    return " ".join(str(name or "").split()).lower()


def money(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    try:
        return Decimal(cleaned or "0")
    except InvalidOperation as exc:
        raise OracleError(f"cannot read {value!r} as money") from exc


def quantity(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise OracleError(f"cannot read {value!r} as a quantity") from exc


def when(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise OracleError(f"cannot read {value!r} as a date")


# ---------------------------------------------------------------------------
# RegisterOne: a database, so the truth is a query
# ---------------------------------------------------------------------------


@dataclass
class SqlOracle:
    """Runs a golden question's SQL against the fake POS's own database."""

    session: AsyncSession

    async def value(self, question: dict[str, Any]) -> Decimal | None:
        statement = question.get("sql")
        if not statement:
            return None
        rows = (await self.session.execute(text(statement))).all()
        if len(rows) != 1 or len(rows[0]) != 1:
            raise OracleError(f"{question['id']}: expected one value, got {len(rows)} rows")
        found = rows[0][0]
        if found is None:
            raise OracleError(f"{question['id']}: the source returned no value")
        return Decimal(str(found))


# ---------------------------------------------------------------------------
# Panel & Pawn: a spreadsheet, so the truth is arithmetic over the raw rows
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class SaleRow:
    receipt: str
    day: date
    item: str
    category: str
    units: Decimal
    net: Decimal


@dataclass
class SpreadsheetOracle:
    """Measures computed straight off the export, by hand.

    The golden file names a measure and a window rather than embedding code, so
    a question stays readable and there is one place where "net sales" means
    quantity times price less discount.
    """

    path: Path
    sales: list[SaleRow] = field(default_factory=list)
    stock: dict[str, dict[str, Decimal]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        for sheet in workbook.sheetnames:
            if sheet.startswith("Sales"):
                self._read_sales(workbook[sheet])
            elif sheet == "Inventory":
                self._read_stock(workbook[sheet])
        workbook.close()
        if not self.sales:
            raise OracleError(f"no sales rows in {self.path}")
        self._settle_categories()

    def _settle_categories(self) -> None:
        """A category belongs to an item, not to a row.

        180 of this file's 1,932 sales rows leave the Category cell blank, and
        every one of those items is filed somewhere else in the file. Asked
        "how much did comics bring in", a shop owner means sales of comics —
        not sales of comics where somebody remembered to type the category. So
        each item takes the first category the file gives it, and every row for
        that item counts under it.
        """
        settled: dict[str, str] = {}
        for row in self.sales:
            if row.category and row.item not in settled:
                settled[row.item] = row.category
        self.sales = [
            replace(row, category=settled.get(row.item, row.category)) for row in self.sales
        ]

    def _read_sales(self, sheet: Any) -> None:
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows)]
        index = {name: position for position, name in enumerate(header)}
        blank_receipts = 0
        for row in rows:
            item = row[index["Item Desc"]]
            if item is None or fold(item) in SUMMARY_ROWS:
                continue
            receipt = str(row[index["Receipt #"]] or "").strip()
            if not receipt:
                # The mapping treats an unnumbered row as its own sale; so does
                # this, or the two disagree on how many orders there were.
                blank_receipts += 1
                receipt = f"__blank_{blank_receipts}"
            units = quantity(row[index["Qty Sold"]])
            price = money(row[index["Unit Price"]])
            discount = money(row[index["Discount"]])
            self.sales.append(
                SaleRow(
                    receipt=receipt,
                    day=when(row[index["Date"]]),
                    item=fold(item),
                    category=fold(row[index["Category"]]),
                    units=units,
                    net=units * price - discount,
                )
            )

    def _read_stock(self, sheet: Any) -> None:
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows)]
        index = {name: position for position, name in enumerate(header)}
        for row in rows:
            item = row[index["Item"]]
            if item is None or fold(item) in SUMMARY_ROWS:
                continue
            self.stock[fold(item)] = {
                "on_hand": quantity(row[index["On Hand"]]),
                "cost": money(row[index["Cost"]]) if row[index["Cost"]] else Decimal("0"),
                "retail": money(row[index["Retail"]]),
                "has_cost": Decimal("1") if row[index["Cost"]] else Decimal("0"),
            }

    def _matching(self, spec: dict[str, Any]) -> list[SaleRow]:
        first = when(spec["from"]) if spec.get("from") else date.min
        last = when(spec["to"]) if spec.get("to") else date.max
        category = fold(spec["category"]) if spec.get("category") else None
        item = fold(spec["item"]) if spec.get("item") else None
        return [
            row
            for row in self.sales
            if first <= row.day <= last
            and (category is None or row.category == category)
            and (item is None or row.item == item)
        ]

    def value(self, question: dict[str, Any]) -> Decimal | None:
        spec = question.get("expect")
        if not spec:
            return None
        measure = spec["measure"]

        if measure in {"on_hand", "retail_value", "cost_value", "items_out_of_stock"}:
            return self._stock_measure(measure)

        rows = self._matching(spec)
        if measure == "net_sales":
            return sum((row.net for row in rows), Decimal("0"))
        if measure == "units":
            return sum((row.units for row in rows), Decimal("0"))
        if measure == "orders":
            return Decimal(len({row.receipt for row in rows}))
        if measure == "average_order_value":
            receipts = {row.receipt for row in rows}
            if not receipts:
                raise OracleError(f"{question['id']}: no sales in that window")
            return sum((row.net for row in rows), Decimal("0")) / Decimal(len(receipts))
        if measure == "top_product_net":
            by_item: dict[str, Decimal] = defaultdict(Decimal)
            for row in rows:
                by_item[row.item] += row.net
            if not by_item:
                raise OracleError(f"{question['id']}: no sales in that window")
            return max(by_item.values())
        raise OracleError(f"{question['id']}: unknown measure {measure!r}")

    def _stock_measure(self, measure: str) -> Decimal:
        if measure == "on_hand":
            return sum((row["on_hand"] for row in self.stock.values()), Decimal("0"))
        if measure == "retail_value":
            return sum(
                (row["on_hand"] * row["retail"] for row in self.stock.values()), Decimal("0")
            )
        if measure == "cost_value":
            return sum((row["on_hand"] * row["cost"] for row in self.stock.values()), Decimal("0"))
        return Decimal(sum(1 for row in self.stock.values() if row["on_hand"] == 0))
