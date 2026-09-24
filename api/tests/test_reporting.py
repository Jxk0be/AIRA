"""The documents that leave the app.

A purchase order goes to a distributor and a month-end packet goes to a
bookkeeper, so both have to be real files that open in whatever those people
use. These tests read the bytes back rather than trusting that a writer wrote.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pypdf

from app.reporting import Column, Document, Sheet, csv_bytes, money, quantity, workbook_bytes


def sample() -> Document:
    doc = Document("Purchase order PO-2026-09-0001", "Tsundoku & Tabletop")
    doc.field("Vendor", "Paper Lantern Books")
    doc.heading("Items")
    doc.table(
        [Column("Item", 260), Column("Qty", 60, "right"), Column("Line", 80, "right")],
        [["Crimson Ronin — Vol. 7", "12", "$85.68"], ["Harborline", "6", "$42.84"]],
        total=["Total", "18", "$128.52"],
    )
    doc.note("Figures at cost where a cost is on file.")
    doc.footer("Tsundoku & Tabletop")
    return doc


def text_of(pdf: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    return "\n".join(page.extract_text() for page in reader.pages)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_the_pdf_is_a_pdf_a_reader_will_open() -> None:
    rendered = sample().render()
    assert rendered.startswith(b"%PDF-1.4")
    assert rendered.rstrip().endswith(b"%%EOF")
    assert len(pypdf.PdfReader(io.BytesIO(rendered)).pages) == 1


def test_every_figure_survives_into_the_text() -> None:
    """A PDF whose numbers are drawn but not extractable is a PDF nobody can
    check against their own books."""
    extracted = text_of(sample().render())
    for expected in ("PO-2026-09-0001", "Crimson Ronin", "$85.68", "18", "$128.52"):
        assert expected in extracted


def test_output_is_byte_identical_for_the_same_input() -> None:
    """No timestamps, no ids, no dictionary ordering — which is what lets a
    packet downloaded in March be the packet generated in February."""
    assert sample().render() == sample().render()


def test_a_long_document_breaks_into_pages() -> None:
    doc = Document("Month end", "Tsundoku & Tabletop")
    doc.table(
        [Column("Item", 300), Column("Net", 100, "right")],
        [[f"Product {index}", f"${index}.00"] for index in range(120)],
    )
    reader = pypdf.PdfReader(io.BytesIO(doc.render()))
    assert len(reader.pages) > 1


def test_characters_the_core_fonts_cannot_show_are_transliterated() -> None:
    """An em dash becoming a hyphen is fine. A mojibake product name is not."""
    doc = Document("Test")
    doc.text("Crimson Ronin — Vol. 7 “special” … 日本語")
    extracted = text_of(doc.render())
    assert "Crimson Ronin - Vol. 7" in extracted
    assert "..." in extracted


def test_money_shows_nothing_rather_than_zero_when_there_is_no_figure() -> None:
    """$0.00 and "we do not know" are different claims and must not print the
    same on something a bookkeeper reads."""
    assert money(None) == "—"
    assert money(Decimal("0")) == "$0.00"
    assert money(Decimal("1234.5")) == "$1,234.50"


def test_quantities_print_the_way_a_person_writes_them() -> None:
    assert quantity(Decimal("12.0000")) == "12"
    assert quantity(Decimal("0.5000")) == "0.5"


# ---------------------------------------------------------------------------
# CSV and XLSX
# ---------------------------------------------------------------------------


def sheet() -> Sheet:
    return Sheet(
        name="Summary",
        note="Figures as recorded in the POS.",
        columns=["figure", "amount"],
        rows=[["Net sales", Decimal("1483.17")], ["Orders", 29], ["Missing", None]],
    )


def test_csv_carries_a_bom_so_excel_does_not_mangle_it() -> None:
    """The overwhelmingly likely next step is a double-click on Windows."""
    assert csv_bytes(sheet()).startswith(b"\xef\xbb\xbf")


def test_csv_keeps_the_note_and_blanks_a_missing_figure() -> None:
    text = csv_bytes(sheet()).decode("utf-8-sig")
    assert text.splitlines()[0].startswith("# Figures as recorded")
    assert "Missing," in text
    assert "1483.17" in text


def test_the_workbook_writes_numbers_as_numbers() -> None:
    """A spreadsheet where every total is text because it came through with a
    dollar sign is worse than no spreadsheet."""
    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes([sheet()])))
    tab = book["Summary"]
    values = {row[0].value: row[1].value for row in tab.iter_rows(min_row=3)}
    assert values["Net sales"] == 1483.17
    assert isinstance(values["Net sales"], float)
    assert values["Orders"] == 29


def test_one_tab_per_sheet_with_a_safe_name() -> None:
    long_name = "By category / location: everything, and then some more"
    book = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes([sheet(), Sheet(name=long_name, columns=["a"], rows=[["b"]])]))
    )
    assert "Summary" in book.sheetnames
    assert len(book.sheetnames) == 2
    other = next(name for name in book.sheetnames if name != "Summary")
    assert len(other) <= 31
    assert "/" not in other
